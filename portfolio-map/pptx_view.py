#!/usr/bin/env python3
"""Render the laid-out page to a one-slide PowerPoint deck.

Every box is a native PowerPoint shape, so the slide can be nudged, recoloured
and copied like any hand-drawn one — but when the content changes you
regenerate it instead of editing it.
"""
from __future__ import annotations

import pathlib

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt

import layout

ALIGN = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr.lstrip("#").upper())


def _fill_text(shape, sh: dict) -> None:
    tf = shape.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE if sh["valign"] == "m" else MSO_ANCHOR.TOP
    for n, ln in enumerate(sh["lines"]):
        para = tf.paragraphs[0] if n == 0 else tf.add_paragraph()
        para.alignment = ALIGN[sh["align"]]
        para.line_spacing = Pt(ln["size"] * 1.28)
        run = para.add_run()
        run.text = ln["t"]
        run.font.size = Pt(ln["size"])
        run.font.bold = ln["bold"]
        run.font.color.rgb = rgb(ln["color"])


ARROW = {"down": MSO_SHAPE.DOWN_ARROW, "right": MSO_SHAPE.RIGHT_ARROW,
         "leftright": MSO_SHAPE.LEFT_RIGHT_ARROW}


def build(pages, path: pathlib.Path) -> pathlib.Path:
    if isinstance(pages, layout.Page):
        pages = [pages]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Pt(layout.W), Pt(layout.H)
    for page in pages:
        _slide(prs, page)
    prs.save(str(path))
    return path


def _slide(prs, page: layout.Page) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    for sh in page.shapes:
        x, y, w, h = (Pt(sh[k]) for k in ("x", "y", "w", "h"))
        if sh["kind"] == "text":
            _fill_text(slide.shapes.add_textbox(x, y, w, h), sh)
            continue

        if sh["kind"] == "ellipse":
            shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, x, y, w, h)
        elif sh["kind"] == "arrow":
            shape = slide.shapes.add_shape(ARROW[sh["dir"]], x, y, w, h)
        elif sh["kind"] == "chevron":
            shape = slide.shapes.add_shape(MSO_SHAPE.CHEVRON, x, y, w, h)
            shape.adjustments[0] = layout.NOTCH / min(sh["w"], sh["h"])
        elif sh["radius"]:
            shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
            shape.adjustments[0] = min(sh["radius"] / min(sh["w"], sh["h"]), 0.5)
        else:
            shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)

        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(sh["fill"])
        shape.shadow.inherit = False
        stroke = sh.get("stroke") if sh["kind"] != "chevron" else "#DEE2E6"
        if stroke:
            shape.line.color.rgb = rgb(stroke)
            shape.line.width = Pt(0.75)
            if sh.get("dash"):
                shape.line.dash_style = 4  # msoLineDash
        else:
            shape.line.fill.background()

        if sh["kind"] == "chevron":
            inset = layout.NOTCH + 8
            _fill_text(shape, {**sh, "align": "l", "valign": "m"})
            shape.text_frame.margin_left = Pt(inset)
        else:
            shape.text_frame.text = ""
